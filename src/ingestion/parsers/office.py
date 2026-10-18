"""Office document extraction: DOCX and PPTX (Task 19b).

Word paragraphs and tables are walked in document order (flow layout has no
page numbers, so blocks carry no page); PowerPoint slides become pages, and
text frames, tables, and grouped shapes flatten in shape order. OOXML
containers pass the shared zip-safety guard before any parsing.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from ingestion.parsers.archives import validate_zip
from ingestion.parsers.basic import ExtractionFailed, RawBlock


def parse_docx(data: bytes) -> list[RawBlock]:
    # Heavy office libs import lazily, matching basic.parse_pdf.
    from docx import Document as load_docx
    from docx.text.paragraph import Paragraph

    validate_zip(data)
    try:
        document = load_docx(BytesIO(data))
    except Exception as exc:
        raise ExtractionFailed("corrupt_docx") from exc

    blocks: list[RawBlock] = []
    for item in document.iter_inner_content():  # yields Paragraph and Table
        if isinstance(item, Paragraph):
            if text := item.text.strip():
                blocks.append(RawBlock(text))
        else:
            _append_rows(item, blocks)
    return blocks


def parse_pptx(data: bytes) -> list[RawBlock]:
    from pptx import Presentation as load_pptx

    validate_zip(data)
    try:
        presentation = load_pptx(BytesIO(data))
    except Exception as exc:
        raise ExtractionFailed("corrupt_pptx") from exc

    blocks: list[RawBlock] = []
    for page, slide in enumerate(presentation.slides, start=1):
        _walk_shapes(slide.shapes, page, blocks)
    return blocks


def _walk_shapes(shapes: Any, page: int, blocks: list[RawBlock]) -> None:
    for shape in shapes:
        if nested := getattr(shape, "shapes", None):  # group shapes
            _walk_shapes(nested, page, blocks)
        elif getattr(shape, "has_table", False):
            _append_rows(shape.table, blocks, page)
        elif text_frame := getattr(shape, "text_frame", None):
            for paragraph in text_frame.paragraphs:
                if text := paragraph.text.strip():
                    blocks.append(RawBlock(text, page))


def _append_rows(table: Any, blocks: list[RawBlock], page: int | None = None) -> None:
    for row in table.rows:
        if text := " | ".join(filter(None, (cell.text.strip() for cell in row.cells))):
            blocks.append(RawBlock(text, page))
