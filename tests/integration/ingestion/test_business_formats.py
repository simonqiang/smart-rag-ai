"""Business document and OCR adapters (Task 19): HTML, DOCX, PPTX, XLSX, CSV, images.

Parser-level tests run in-memory; upload acceptance runs against the live
Compose database through the shared ingestion fixtures.
"""

import asyncio
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from sqlalchemy.ext.asyncio import create_async_engine

from foundation.events import EventWriter
from foundation.storage import ObjectStore
from foundation.unit_of_work import UnitOfWork
from identity_access.authorization import AccessContext
from ingestion.extraction import SourceObject, extract
from ingestion.parsers.archives import MAX_DECOMPRESSED_BYTES
from ingestion.parsers.basic import ExtractionFailed

IMAGE_MEDIA = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "webp": "image/webp",
    "tif": "image/tiff",
    "bmp": "image/bmp",
}

DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MEDIA = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _extract(media_type: str, data: bytes, filename: str) -> object:
    return extract(
        SourceObject(
            source_version_id="v1", media_type=media_type, filename=filename, data=data
        )
    )


def _html(data: bytes) -> object:
    return _extract("text/html", data, "page.html")


# --------------------------------------------------------------------------
# HTML (19a)
# --------------------------------------------------------------------------


def test_html_scripts_styles_and_templates_are_dropped() -> None:
    doc = _html(
        b"<html><head><script>alert(1)</script><style>p { color: red }</style>"
        b"</head><body><template><p>ghost</p></template>"
        b"<h1>Policy</h1><p>Leave accrues monthly.</p>"
        b"<noscript>enable javascript</noscript></body></html>"
    )

    assert [block.text for block in doc.blocks] == ["Policy", "Leave accrues monthly."]
    assert doc.warnings == []


def test_html_nav_header_footer_aside_boilerplate_is_skipped() -> None:
    doc = _html(
        b"<html><body><header>ACME Inc</header>"
        b"<nav><header>brand</header><ul><li>Home</li><li>Products</li></ul></nav>"
        b"<article><p>Real content stays.</p></article>"
        b"<aside>Advertisement</aside><footer>Copyright ACME</footer></body></html>"
    )

    assert [block.text for block in doc.blocks] == ["Real content stays."]


def test_html_blocks_are_ordered_with_ordinal_locations() -> None:
    doc = _html(
        b"<html><body><h1>Handbook</h1>"
        b"<p>First paragraph across<br>two lines.</p>"
        b"<ul><li>Alpha</li><li>Beta</li></ul></body></html>"
    )

    assert [block.text for block in doc.blocks] == [
        "Handbook",
        "First paragraph across two lines.",
        "Alpha",
        "Beta",
    ]
    assert [block.location.block for block in doc.blocks] == [0, 1, 2, 3]
    assert all(block.location.page is None for block in doc.blocks)


def test_html_gb18030_meta_charset_decodes_chinese() -> None:
    body = (
        '<html><head><meta charset="gb18030"></head>'
        "<body><h1>公司介绍</h1><p>员工每年享有十五天年假。</p></body></html>"
    ).encode("gb18030")
    doc = _html(body)

    assert [block.text for block in doc.blocks] == [
        "公司介绍",
        "员工每年享有十五天年假。",
    ]
    assert doc.blocks[-1].language == "zh-Hans"


def test_html_undeclared_utf8_chinese_still_decodes() -> None:
    doc = _html("<html><body><p>員工享有年假。</p></body></html>".encode())

    assert doc.blocks[0].text == "員工享有年假。"
    assert doc.blocks[0].language == "zh-Hant"


def test_html_malformed_markup_still_extracts_readable_text() -> None:
    doc = _html(
        b"<html><body></script><p>Unclosed <b>bold and <i>more text</p><p>Next</b></p>"
    )

    assert [block.text for block in doc.blocks] == [
        "Unclosed bold and more text",
        "Next",
    ]


def test_html_undecodable_bytes_fail_typed() -> None:
    from ingestion.parsers.basic import ExtractionFailed

    with pytest.raises(ExtractionFailed) as excinfo:
        _html(b"\xff\xfe\xff\x00\x81\x81 raw junk")
    assert excinfo.value.reason == "decode_failed"


def test_html_script_only_page_warns_empty() -> None:
    doc = _html(
        b"<html><head><title></title></head><body><script>var x=1;</script></body></html>"
    )

    assert doc.blocks == []
    assert doc.warnings == ["empty"]


def test_html_entities_and_comments_render_as_text() -> None:
    doc = _html(
        b"<html><body><p>Profit &amp; loss &lt;2026&gt;</p><!-- hidden --></body></html>"
    )

    assert doc.blocks[0].text == "Profit & loss <2026>"


def test_html_wrong_declared_charset_falls_back_to_detection() -> None:
    content = "<html><body><p>員工享有年假。</p></body></html>".encode()  # really UTF-8
    doc = _html(content.replace(b"<body>", b'<body><meta charset="iso-8859-1">'))
    doc_late_meta = _html(
        content.replace(
            b"<body>",
            b'<body><meta http-equiv="content-type" content="text/html; charset=gb18030">',
        )
    )

    assert doc.blocks[0].text == "員工享有年假。"
    assert doc_late_meta.blocks[0].text == "員工享有年假。"
    assert doc.blocks[0].language == "zh-Hant"


def test_html_blank_block_content_is_dropped() -> None:
    doc = _html(b"<html><body><p>Real</p><p>   </p><p>Also real</p></body></html>")

    assert [block.text for block in doc.blocks] == ["Real", "Also real"]


def test_html_registered_in_extraction_registry() -> None:
    doc = _extract("text/html", b"<html><body><p>hello</p></body></html>", "page.html")

    assert doc.media_type == "text/html"
    assert doc.blocks[0].text == "hello"


def test_html_upload_accepted_and_stored(db, store_root, workspace) -> None:
    from ingestion.uploads import register_upload

    async def run(uow, store: ObjectStore) -> object:
        return await register_upload(
            uow,
            EventWriter(),
            store,
            context=AccessContext(
                user_id=workspace["owner_id"],
                workspace_id=workspace["workspace_id"],
                role="owner",
            ),
            collection_id=workspace["collection_id"],
            name="Landing page",
            filename="page.html",
            data=b"<html><body><h1>Hello</h1></body></html>",
        )

    async def run_all() -> object:
        engine = create_async_engine(db.database_url)
        uow = UnitOfWork(engine)
        try:
            return await run(uow, ObjectStore(store_root))
        finally:
            await engine.dispose()

    accepted = asyncio.run(run_all())

    assert accepted.media_type == "text/html"


# --------------------------------------------------------------------------
# DOCX and PPTX (19b)
# --------------------------------------------------------------------------


def _docx(*paragraphs: tuple[str, str]) -> bytes:
    from docx import Document

    document = Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_slides(*slides: list[tuple[str, str]]) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    for shapes in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        for placeholder, text in shapes:
            slide.placeholders[1 if placeholder == "body" else 0].text = text
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_docx_headings_and_paragraphs_ordered_multilingual() -> None:
    doc = _extract(
        DOCX_MEDIA,
        _docx(
            ("Heading 1", "Quarterly Report"),
            ("Normal", "Revenue grew 12 percent."),
            ("Normal", "营收增长百分之十二。"),
        ),
        "report.docx",
    )

    assert [block.text for block in doc.blocks] == [
        "Quarterly Report",
        "Revenue grew 12 percent.",
        "营收增长百分之十二。",
    ]
    assert all(block.location.page is None for block in doc.blocks)
    assert doc.blocks[-1].language == "zh-Hans"


def test_docx_tables_extracted_in_document_order() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Price"
    document.add_paragraph("After table")
    buffer = BytesIO()
    document.save(buffer)
    doc = _extract(DOCX_MEDIA, buffer.getvalue(), "prices.docx")

    assert [block.text for block in doc.blocks] == [
        "Before table",
        "Item | Price",
        "After table",
    ]


def test_docx_empty_document_warns_empty() -> None:
    doc = _extract(DOCX_MEDIA, _docx(), "empty.docx")

    assert doc.blocks == []
    assert doc.warnings == ["empty"]


def test_docx_corrupt_bytes_fail_typed() -> None:
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(DOCX_MEDIA, b"definitely not a zip file", "broken.docx")
    assert excinfo.value.reason == "corrupt_archive"


def test_docx_valid_zip_with_garbage_content_fails_typed() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", b"<not-a-word-document/>")
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(DOCX_MEDIA, buffer.getvalue(), "junk.docx")
    assert excinfo.value.reason == "corrupt_docx"


def test_docx_encrypted_zip_fails_typed() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", b"<xml/>")

    def encrypted(data: bytes) -> bytes:
        # zipfile refuses to write encrypted entries; flip the central-directory
        # flag bit the way a real encrypted archive carries it.
        out = bytearray(data)
        header = out.index(b"PK\x01\x02")
        flags = int.from_bytes(out[header + 8 : header + 10], "little") | 0x1
        out[header + 8 : header + 10] = flags.to_bytes(2, "little")
        return bytes(out)

    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(DOCX_MEDIA, encrypted(buffer.getvalue()), "secret.docx")
    assert excinfo.value.reason == "encrypted_archive"


def test_docx_archive_bomb_fails_typed() -> None:
    bomb = BytesIO()
    with ZipFile(bomb, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (MAX_DECOMPRESSED_BYTES + 1))
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(DOCX_MEDIA, bomb.getvalue(), "bomb.docx")
    assert excinfo.value.reason == "archive_bomb"


def test_pptx_slides_carry_page_locations() -> None:
    doc = _extract(
        PPTX_MEDIA,
        _pptx_slides(
            [("title", "Strategy"), ("body", " invest in APAC ")],
            [("title", "Roadmap"), ("body", "Q4 launch")],
        ),
        "deck.pptx",
    )

    assert [(block.text, block.location.page) for block in doc.blocks] == [
        ("Strategy", 1),
        ("invest in APAC", 1),
        ("Roadmap", 2),
        ("Q4 launch", 2),
    ]


def test_pptx_tables_and_grouped_shapes_extracted() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    table_shape = slide.shapes.add_table(
        1, 2, Inches(1), Inches(1), Inches(4), Inches(1)
    )
    table_shape.table.cell(0, 0).text = "Region"
    table_shape.table.cell(0, 1).text = "Sales"
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(1), Inches(1)
    ).text_frame.text = "Note"
    buffer = BytesIO()
    presentation.save(buffer)
    doc = _extract(PPTX_MEDIA, buffer.getvalue(), "grid.pptx")

    assert [(block.text, block.location.page) for block in doc.blocks] == [
        ("Region | Sales", 1),
        ("Note", 1),
    ]


def test_pptx_corrupt_bytes_fail_typed() -> None:
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(PPTX_MEDIA, b"\x00\x01\x02 not a deck", "broken.pptx")
    assert excinfo.value.reason == "corrupt_archive"

    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", b"<not-a-deck/>")
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(PPTX_MEDIA, buffer.getvalue(), "junk.pptx")
    assert excinfo.value.reason == "corrupt_pptx"


def test_office_uploads_reject_non_zip_content(db, store_root, workspace) -> None:
    from ingestion.uploads import UploadRejected, register_upload

    async def run(uow, store: ObjectStore) -> object:
        return await register_upload(
            uow,
            EventWriter(),
            store,
            context=AccessContext(
                user_id=workspace["owner_id"],
                workspace_id=workspace["workspace_id"],
                role="owner",
            ),
            collection_id=workspace["collection_id"],
            name="Broken deck",
            filename="broken.pptx",
            data=b"plain text masquerading as a deck",
        )

    async def run_all() -> object:
        engine = create_async_engine(db.database_url)
        uow = UnitOfWork(engine)
        try:
            return await run(uow, ObjectStore(store_root))
        finally:
            await engine.dispose()

    with pytest.raises(UploadRejected) as excinfo:
        asyncio.run(run_all())
    assert excinfo.value.reason == "corrupt_office"


def test_docx_table_only_document_extracts_rows() -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Widget"
    table.cell(1, 1).text = "5"
    buffer = BytesIO()
    document.save(buffer)
    doc = _extract(DOCX_MEDIA, buffer.getvalue(), "table.docx")

    assert [block.text for block in doc.blocks] == ["Item | Price", "Widget | 5"]


def test_docx_blank_rows_and_cells_are_skipped() -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = " "
    table.cell(1, 0).text = "Widget"
    document.add_paragraph("")  # blank paragraph between rows and tail
    document.add_paragraph("Tail")
    buffer = BytesIO()
    document.save(buffer)
    doc = _extract(DOCX_MEDIA, buffer.getvalue(), "sparse.docx")

    assert [block.text for block in doc.blocks] == ["Widget", "Tail"]


def test_pptx_pictures_and_blank_paragraphs_are_skipped() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    slide.shapes.add_picture(BytesIO(_png_pixel()), Inches(0), Inches(0))
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text_frame.text = "Visible"
    box.text_frame.add_paragraph()  # trailing blank paragraph
    buffer = BytesIO()
    presentation.save(buffer)
    doc = _extract(PPTX_MEDIA, buffer.getvalue(), "mixed.pptx")

    assert [(block.text, block.location.page) for block in doc.blocks] == [
        ("Visible", 1)
    ]


def _png_pixel() -> bytes:
    # A 1x1 red PNG; keeps the picture-shape test independent of Pillow.
    return bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753"
        "de0000000c4944415408d763f8cfc000000301010018dd8db00000000049454e44ae426082"
    )


def test_zip_guard_refuses_more_than_ten_thousand_entries() -> None:
    from ingestion.parsers.archives import MAX_ENTRIES

    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for index in range(MAX_ENTRIES + 1):
            archive.writestr(f"part/{index}.xml", b"<x/>")
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(DOCX_MEDIA, buffer.getvalue(), "many.docx")
    assert excinfo.value.reason == "archive_bomb"


# --------------------------------------------------------------------------
# XLSX and CSV (19c)
# --------------------------------------------------------------------------


def _xlsx(*sheets: list[list[object]], merged: bool = False) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for rows in sheets:
        sheet = workbook.create_sheet()
        for row in rows:
            sheet.append(row)
    if merged:
        workbook.worksheets[0].merge_cells("A2:B2")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_sheets_are_pages_with_header_paired_rows() -> None:
    doc = _extract(
        XLSX_MEDIA,
        _xlsx(
            [["Item", "Price"], ["Widget", 5], ["Gadget", 12.5]],
            [["Region", "Q1"], ["APAC", 100]],
        ),
        "prices.xlsx",
    )

    assert [(block.text, block.location.page) for block in doc.blocks] == [
        ("Item=Widget | Price=5", 1),
        ("Item=Gadget | Price=12.5", 1),
        ("Region=APAC | Q1=100", 2),
    ]


def test_xlsx_merged_cells_keep_top_left_value_only() -> None:
    doc = _extract(
        XLSX_MEDIA,
        _xlsx([["Item", "Price"], ["Bundle deal", None], ["Widget", 5]], merged=True),
        "merged.xlsx",
    )

    assert [block.text for block in doc.blocks] == [
        "Item=Bundle deal",
        "Item=Widget | Price=5",
    ]


def test_xlsx_uncached_formula_contributes_nothing() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Total"])
    sheet["A2"] = 2
    sheet["A3"] = 3
    sheet["A4"] = "=A2+A3"  # cached value written by openpyxl on save? (no)
    buffer = BytesIO()
    workbook.save(buffer)
    doc = _extract(XLSX_MEDIA, buffer.getvalue(), "formulas.xlsx")

    # plain values pass through; openpyxl never writes caches, so the
    # uncached =A2+A3 formula contributes nothing.
    assert [block.text for block in doc.blocks] == ["Total=2", "Total=3"]


def test_xlsx_empty_sheets_warn_empty() -> None:
    doc = _extract(XLSX_MEDIA, _xlsx([[], []]), "blank.xlsx")

    assert doc.blocks == []
    assert doc.warnings == ["empty"]


def test_xlsx_large_sheet_is_truncated_with_warning() -> None:
    from ingestion.parsers.spreadsheet import MAX_SHEET_ROWS

    rows: list[list[object]] = [["Item", "Price"]]
    rows += [["Widget", index] for index in range(MAX_SHEET_ROWS + 5)]
    doc = _extract(XLSX_MEDIA, _xlsx(rows), "large.xlsx")

    assert len(doc.blocks) == MAX_SHEET_ROWS
    assert doc.warnings == ["rows_truncated"]


def test_xlsx_overlong_cell_is_truncated_with_warning() -> None:
    from ingestion.parsers.spreadsheet import MAX_CELL_CHARS

    doc = _extract(
        XLSX_MEDIA, _xlsx([["Item"], ["x" * (MAX_CELL_CHARS + 50)]]), "long.xlsx"
    )

    assert len(doc.blocks[0].text) < MAX_CELL_CHARS + 10
    assert doc.warnings == ["cells_truncated"]


def test_xlsx_corrupt_zip_fails_typed() -> None:
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(XLSX_MEDIA, b"PK\x03\x04 followed by junk", "broken.xlsx")
    assert excinfo.value.reason == "corrupt_archive"

    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", b"<not-a-workbook/>")
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract(XLSX_MEDIA, buffer.getvalue(), "junk.xlsx")
    assert excinfo.value.reason == "corrupt_xlsx"


def test_csv_rows_are_header_paired_without_pages() -> None:
    doc = _extract("text/csv", b"Item,Price\nWidget,5\nGadget,12.5\n", "prices.csv")

    assert [(block.text, block.location.page) for block in doc.blocks] == [
        ("Item=Widget | Price=5", None),
        ("Item=Gadget | Price=12.5", None),
    ]


def test_csv_gb18030_decodes_chinese_headers() -> None:
    body = "项目,价格\nwidget,5\n".encode("gb18030")
    doc = _extract("text/csv", body, "prices.csv")

    assert doc.blocks[0].text == "项目=widget | 价格=5"
    assert doc.blocks[0].language == "mixed"  # CJK headers + latin values


def test_csv_blank_rows_and_extra_columns_are_kept_positional() -> None:
    doc = _extract("text/csv", b"Item,Price\nWidget,5,extra\n,7\n", "ragged.csv")

    assert [block.text for block in doc.blocks] == [
        "Item=Widget | Price=5 | Column3=extra",
        "Price=7",
    ]


def test_csv_nul_bytes_are_stripped() -> None:
    doc = _extract("text/csv", b"Item,Price\nW\x00i,d\n", "null.csv")

    assert [block.text for block in doc.blocks] == ["Item=Wi | Price=d"]


def test_csv_overgrown_field_fails_typed() -> None:
    from ingestion.parsers.spreadsheet import MAX_FIELD_CHARS

    with pytest.raises(ExtractionFailed) as excinfo:
        _extract("text/csv", b"A\n" + b"x" * (MAX_FIELD_CHARS + 1), "huge.csv")
    assert excinfo.value.reason == "corrupt_csv"


def test_csv_large_file_is_truncated_with_warning() -> None:
    from ingestion.parsers.spreadsheet import MAX_SHEET_ROWS

    body = b"Item,Price\n" + b"".join(
        b"Widget,%d\n" % index for index in range(MAX_SHEET_ROWS + 5)
    )
    doc = _extract("text/csv", body, "large.csv")

    assert len(doc.blocks) == MAX_SHEET_ROWS
    assert doc.warnings == ["rows_truncated"]


def test_xlsx_and_csv_uploads_accepted(db, store_root, workspace) -> None:
    from ingestion.uploads import register_upload

    async def run(uow, store: ObjectStore, filename: str, data: bytes) -> object:
        return await register_upload(
            uow,
            EventWriter(),
            store,
            context=AccessContext(
                user_id=workspace["owner_id"],
                workspace_id=workspace["workspace_id"],
                role="owner",
            ),
            collection_id=workspace["collection_id"],
            name="Sheet",
            filename=filename,
            data=data,
        )

    async def run_all() -> list[object]:
        engine = create_async_engine(db.database_url)
        uow = UnitOfWork(engine)
        store = ObjectStore(store_root)
        try:
            return [
                await run(uow, store, "sheet.xlsx", _xlsx([["A"], [1]])),
                await run(uow, store, "sheet.csv", b"A,B\n1,2\n"),
            ]
        finally:
            await engine.dispose()

    accepted = asyncio.run(run_all())

    assert [upload.media_type for upload in accepted] == [
        XLSX_MEDIA,
        "text/csv",
    ]


# --------------------------------------------------------------------------
# Image OCR (19d) — requires the host Tesseract with eng/chi_sim/chi_tra packs
# --------------------------------------------------------------------------


OCR_FIXTURES = Path(__file__).resolve().parents[3] / "fixtures" / "ocr"


def _image(media_type: str, filename: str) -> object:
    return _extract(media_type, (OCR_FIXTURES / filename).read_bytes(), filename)


def test_ocr_extracts_english_text_on_one_page() -> None:
    doc = _image("image/png", "en.png")

    assert doc.warnings == []
    text = " ".join(block.text for block in doc.blocks)
    assert "Quarterly Revenue Report" in text
    assert doc.blocks[0].location.page == 1
    assert doc.blocks[0].language == "en"


def test_ocr_extracts_malay_text() -> None:
    doc = _image("image/png", "ms.png")

    text = " ".join(block.text for block in doc.blocks)
    assert "Cuti tahunan" in text
    assert doc.blocks[0].language == "ms"


def test_ocr_extracts_simplified_chinese_content() -> None:
    doc = _image("image/png", "zh-hans.png")

    # Script-neutral substring: identical in both scripts regardless of which
    # pack wins the confidence arbitration.
    text = " ".join(block.text for block in doc.blocks)
    assert "每年享有十五天" in text


def test_ocr_extracts_traditional_chinese_content() -> None:
    doc = _image("image/png", "zh-hant.png")

    text = " ".join(block.text for block in doc.blocks)
    assert "享有十五天" in text


def _tsv(rows: list[tuple[int, int, str, float]]) -> dict:
    return {
        "block_num": [row[0] for row in rows],
        "par_num": [row[1] for row in rows],
        "text": [row[2] for row in rows],
        "conf": [row[3] for row in rows],
    }


def test_ocr_arbitration_prefers_higher_confidence_pass(monkeypatch) -> None:
    import pytesseract

    from ingestion.parsers import ocr

    calls: list[str] = []

    def fake_tsv(image, lang: str, **kwargs) -> dict:
        calls.append(lang)
        if lang == ocr.BASE_LANG:
            return _tsv([(1, 1, "員工", 90.0)])
        return _tsv([(1, 1, "员工", 95.0)])

    monkeypatch.setattr(ocr, "_load", lambda data: object())
    monkeypatch.setattr(pytesseract, "image_to_data", fake_tsv)
    doc = _extract("image/png", b"unused", "scan.png")

    assert calls == [ocr.BASE_LANG, ocr.ALT_LANG]
    assert doc.blocks[0].text == "员工"


def test_ocr_arbitration_keeps_base_pass_on_tie(monkeypatch) -> None:
    import pytesseract

    from ingestion.parsers import ocr

    monkeypatch.setattr(ocr, "_load", lambda data: object())
    monkeypatch.setattr(
        pytesseract,
        "image_to_data",
        lambda image, lang, **kwargs: _tsv([(1, 1, "员工年假", 90.0)]),
    )
    doc = _extract("image/png", b"unused", "scan.png")

    assert doc.blocks[0].text == "员工年假"  # tie prefers chi_sim


def test_ocr_groups_tsv_words_into_paragraph_blocks(monkeypatch) -> None:
    import pytesseract

    from ingestion.parsers import ocr

    monkeypatch.setattr(ocr, "_load", lambda data: object())
    monkeypatch.setattr(
        pytesseract,
        "image_to_data",
        lambda image, lang, **kwargs: _tsv(
            [
                (1, 1, "Cuti", 90.0),
                (1, 1, "tahunan", 92.0),
                (1, 1, "", -1.0),  # non-word row
                (1, 1, "?", "-"),  # unparseable confidence reads as no-text
                (1, 2, "untuk", 88.0),
                (2, 1, "kakitangan", 95.0),
            ]
        ),
    )
    doc = _extract("image/png", b"unused", "scan.png")

    assert [block.text for block in doc.blocks] == [
        "Cuti tahunan",
        "untuk",
        "kakitangan",
    ]
    assert [block.location.page for block in doc.blocks] == [1, 1, 1]


def test_ocr_single_pass_for_latin_images(monkeypatch) -> None:
    import pytesseract

    from ingestion.parsers import ocr

    calls: list[str] = []
    monkeypatch.setattr(ocr, "_load", lambda data: object())

    def fake_tsv(image, lang: str, **kwargs) -> dict:
        calls.append(lang)
        return _tsv([(1, 1, "hello", 99.0)])

    monkeypatch.setattr(pytesseract, "image_to_data", fake_tsv)
    doc = _extract("image/png", b"unused", "scan.png")

    assert calls == [ocr.BASE_LANG]  # no Han characters, no second pass
    assert doc.blocks[0].text == "hello"


def test_ocr_maps_tesseract_failures_to_typed_reasons(monkeypatch) -> None:
    import pytesseract

    from ingestion.parsers import ocr

    def raises(exc: Exception):
        def boom(image, lang: str, **kwargs):
            raise exc

        return boom

    cases = [
        (raises(pytesseract.TesseractNotFoundError()), "ocr_unavailable"),
        (raises(pytesseract.TesseractError(-1, "boom")), "ocr_failed"),
        (raises(RuntimeError("Tesseract process timeout")), "ocr_failed"),
    ]
    monkeypatch.setattr(ocr, "_load", lambda data: object())
    for fake, expected in cases:
        monkeypatch.setattr(pytesseract, "image_to_data", fake)
        with pytest.raises(ExtractionFailed) as excinfo:
            _extract("image/png", b"unused", "scan.png")
        assert excinfo.value.reason == expected


def test_ocr_missing_language_pack_is_actionable(monkeypatch) -> None:
    import pytesseract

    from ingestion.parsers import ocr

    def boom(image, lang: str, **kwargs):
        raise pytesseract.TesseractError(-1, "Failed loading language 'chi_tra'")

    monkeypatch.setattr(ocr, "_load", lambda data: object())
    monkeypatch.setattr(pytesseract, "image_to_data", boom)
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract("image/png", b"unused", "scan.png")
    assert excinfo.value.reason == "ocr_language_missing"


def test_ocr_rejects_corrupt_and_oversized_images(monkeypatch) -> None:
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract("image/png", b"this is not an image at all", "junk.png")
    assert excinfo.value.reason == "corrupt_image"

    from ingestion.parsers import ocr

    monkeypatch.setattr(ocr, "MAX_IMAGE_PIXELS", 100)
    with pytest.raises(ExtractionFailed) as excinfo:
        _extract("image/png", (OCR_FIXTURES / "en.png").read_bytes(), "en.png")
    assert excinfo.value.reason == "image_too_large"


def test_image_upload_magic_rejections_and_acceptance() -> None:
    from PIL import Image

    from ingestion.uploads import UploadRejected, _validate

    rejected = [
        ("x.png", b"plain text pretending"),
        ("x.jpg", b"\xff\xd9"),  # wrong JPEG marker
        ("x.webp", b"RIFFxxxxGAME"),  # RIFF container without the WEBP stamp
        ("x.webp", b"GAMExxxxWEBP"),  # WEBP stamp outside a RIFF container
        ("x.tif", b"MM\x00"),  # truncated big-endian magic
        ("x.bmp", b"BA"),  # OS/2 bitmap array, not a supported BMP
    ]
    for filename, data in rejected:
        with pytest.raises(UploadRejected) as excinfo:
            _validate(data, filename, 1000)
        assert excinfo.value.reason == "corrupt_image"

    pixel = Image.new("RGB", (1, 1), "white")
    for fmt, ext in [
        ("PNG", "png"),
        ("JPEG", "jpg"),
        ("WEBP", "webp"),
        ("TIFF", "tif"),
        ("BMP", "bmp"),
    ]:
        buffer = BytesIO()
        pixel.save(buffer, format=fmt)
        assert _validate(buffer.getvalue(), f"scan.{ext}", 1000) == IMAGE_MEDIA[ext]


def test_ocr_flattens_transparent_images_onto_white(monkeypatch) -> None:
    import pytesseract
    from PIL import Image


    image = Image.new("RGBA", (4, 4), (0, 0, 0, 0))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    seen: list[str] = []

    def fake_tsv(img, lang: str, **kwargs) -> dict:
        seen.append((img.mode, img.size))
        return _tsv([(1, 1, "text", 99.0)])

    monkeypatch.setattr(pytesseract, "image_to_data", fake_tsv)
    doc = _extract("image/png", buffer.getvalue(), "transparent.png")

    assert seen == [("RGB", (4, 4))]  # alpha flattened, white background
    assert doc.blocks[0].text == "text"
